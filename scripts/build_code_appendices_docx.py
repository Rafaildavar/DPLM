from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Приложения_код_к_диплому_GestureBind.docx"


SNIPPETS = [
    {
        "appendix": "Приложение А",
        "title": "ORM-модель базы данных",
        "source": "app/models/database.py",
        "parts": [
            ("Листинг А.1 - ORM-модели пользовательского словаря и сэмплов", 106, 203),
            ("Листинг А.2 - ORM-модели обученной модели и команды", 206, 288),
            ("Листинг А.3 - ORM-модели журналов распознавания и сессий", 338, 409),
        ],
    },
    {
        "appendix": "Приложение Б",
        "title": "Сервис записи и синхронизации обучающих примеров",
        "source": "app/services/gesture_samples.py",
        "parts": [
            ("Листинг Б.1 - Создание или обновление записи жеста", 27, 60),
            ("Листинг Б.2 - Запись метаданных обучающего примера в БД", 63, 133),
            ("Листинг Б.3 - Загрузка сэмплов из БД для обучения", 136, 220),
        ],
    },
    {
        "appendix": "Приложение В",
        "title": "Онлайн-распознавание жестов",
        "source": "app/gesture_online_infer.py",
        "parts": [
            ("Листинг В.1 - Инициализация модели и MediaPipe Hand Landmarker", 31, 133),
            ("Листинг В.2 - Обработка кадра и формирование предсказания", 277, 361),
        ],
    },
    {
        "appendix": "Приложение Г",
        "title": "Привязка жестов к командам",
        "source": "app/services/gesture_command_bridge.py",
        "parts": [
            ("Листинг Г.1 - Сохранение пары жест-команда в БД", 193, 255),
            ("Листинг Г.2 - Поиск и выполнение команды по распознанному жесту", 301, 455),
        ],
    },
    {
        "appendix": "Приложение Д",
        "title": "Экран обучения в пользовательском и developer-режимах",
        "source": "app/flet_app/views/training.py",
        "parts": [
            ("Листинг Д.1 - Поля обычного режима и режима разработчика", 30, 77),
            ("Листинг Д.2 - Отображение и удаление записанных сэмплов", 182, 276),
            ("Листинг Д.3 - Запуск записи и обучения модели", 326, 404),
        ],
    },
    {
        "appendix": "Приложение Е",
        "title": "Расчет метрик классификации",
        "source": "scripts/evaluate_gesture_metrics.py",
        "parts": [
            ("Листинг Е.1 - Загрузка датасета sample_*.npy", 61, 101),
            ("Листинг Е.2 - Cross-validation и KNN-предсказание", 104, 142),
            ("Листинг Е.3 - Расчет accuracy, precision, recall и F1", 145, 202),
            ("Листинг Е.4 - Формирование итогового JSON-отчета", 216, 240),
        ],
    },
]


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0


def add_appendix_title(doc: Document, appendix: str, title: str) -> None:
    if len(doc.paragraphs) > 1:
        doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{appendix}. {title}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)


def add_listing_title(doc: Document, title: str, source: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = src.add_run(f"Файл: {source}")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)


def add_code_line(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(7.5)


def add_code_block(doc: Document, source: Path, start: int, end: int) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    for number in range(start, end + 1):
        if number <= 0 or number > len(lines):
            continue
        raw = lines[number - 1].replace("\t", "    ")
        add_code_line(doc, raw)


def main() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("ПРИЛОЖЕНИЯ")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Листинги программного кода GestureBind")
    subtitle_run.italic = True
    subtitle_run.font.name = "Times New Roman"
    subtitle_run.font.size = Pt(12)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.first_line_indent = Inches(0.49)
    note.add_run(
        "В приложениях приведены основные фрагменты программного кода, подтверждающие реализацию базы данных, "
        "записи обучающих примеров, онлайн-распознавания, привязки жестов к командам, пользовательского интерфейса "
        "обучения и расчета метрик классификации."
    )

    for item in SNIPPETS:
        add_appendix_title(doc, item["appendix"], item["title"])
        source = item["source"]
        source_path = ROOT / source
        for idx, (listing_title, start, end) in enumerate(item["parts"]):
            if idx > 0:
                doc.add_paragraph()
            add_listing_title(doc, listing_title, source)
            add_code_block(doc, source_path, start, end)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
