from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "articles" / "drafts" / "Статья_Давар_жесты.docx"
ASSET_DIR = ROOT / "docs" / "articles" / "assets" / "article_assets"


def font_path():
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
        "/System/Library/Fonts/Supplemental/DejaVu Sans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return None


def pil_font(size, bold=False):
    path = font_path()
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines = []
    line = ""
    for word in words:
        trial = (line + " " + word).strip()
        bbox = draw.textbbox((0, 0), trial, font=font)
        if bbox[2] - bbox[0] <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_box(draw, xy, text, fill, outline="#455A64", font=None, title=False):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    font = font or pil_font(30)
    lines = wrap_text(draw, text, font, x2 - x1 - 36)
    line_height = int(font.size * 1.25)
    total = line_height * len(lines)
    y = y1 + ((y2 - y1) - total) // 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, y), line, font=font, fill="#1F2933")
        y += line_height


def arrow(draw, start, end, color="#546E7A"):
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill=color, width=4)
    if ex >= sx:
        points = [(ex, ey), (ex - 16, ey - 9), (ex - 16, ey + 9)]
    else:
        points = [(ex, ey), (ex + 16, ey - 9), (ex + 16, ey + 9)]
    draw.polygon(points, fill=color)


def make_architecture_image(path):
    w, h = 1800, 920
    img = Image.new("RGB", (w, h), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(38)
    box_font = pil_font(30)
    small_font = pil_font(26)
    draw.text((70, 45), "Общая архитектура системы распознавания жестов", font=title_font, fill="#22313F")

    top = 150
    boxes = [
        (80, top, 360, top + 150, "Камера и видеопоток", "#EAF3F8"),
        (450, top, 780, top + 150, "Выделение ключевых точек кисти", "#E9F6EF"),
        (870, top, 1200, top + 150, "Формирование признаков", "#FFF5E5"),
        (1290, top, 1620, top + 150, "ML-модель распознавания", "#F1ECFA"),
    ]
    for x1, y1, x2, y2, text, fill in boxes:
        draw_box(draw, (x1, y1, x2, y2), text, fill, font=box_font)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][2] + 20, top + 75), (boxes[i + 1][0] - 20, top + 75))

    lower = [
        (220, 520, 560, 690, "Словарь пользовательских жестов", "#F7F9FB"),
        (720, 520, 1060, 690, "Обучение и обновление классификатора", "#F7F9FB"),
        (1220, 520, 1560, 690, "Интерфейс и выполнение команд", "#F7F9FB"),
    ]
    for x1, y1, x2, y2, text, fill in lower:
        draw_box(draw, (x1, y1, x2, y2), text, fill, outline="#78909C", font=small_font)

    arrow(draw, (1040, 330), (390, 510))
    arrow(draw, (1040, 330), (890, 510))
    arrow(draw, (1455, 330), (1390, 510))
    arrow(draw, (560, 605), (720, 605))
    arrow(draw, (1060, 605), (1220, 605))

    draw.text((70, 820), "Модульность позволяет изменять интерфейс, классификатор или набор команд без перестройки всего приложения.", font=small_font, fill="#455A64")
    img.save(path)


def make_flow_image(path):
    w, h = 1800, 820
    img = Image.new("RGB", (w, h), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(38)
    box_font = pil_font(29)
    small_font = pil_font(24)
    draw.text((70, 45), "Поток данных при обучении и распознавании", font=title_font, fill="#22313F")

    steps = [
        ("1. Кадр с камеры", "#EAF3F8"),
        ("2. 21 ключевая точка кисти", "#E9F6EF"),
        ("3. Нормализация координат", "#FFF5E5"),
        ("4. Вектор признаков", "#F1ECFA"),
        ("5. Классификация жеста", "#FCEEEF"),
        ("6. Команда в системе", "#EDF7ED"),
    ]
    x, y = 70, 180
    box_w, box_h, gap = 250, 135, 38
    centers = []
    for idx, (text, fill) in enumerate(steps):
        x1 = x + idx * (box_w + gap)
        draw_box(draw, (x1, y, x1 + box_w, y + box_h), text, fill, font=box_font)
        centers.append((x1 + box_w // 2, y + box_h // 2))
        if idx:
            arrow(draw, (x1 - gap + 10, y + box_h // 2), (x1 - 10, y + box_h // 2))

    train_y = 500
    draw.rounded_rectangle((230, train_y, 820, train_y + 170), radius=18, fill="#F7F9FB", outline="#78909C", width=2)
    draw.text((270, train_y + 30), "Режим обучения", font=box_font, fill="#1F2933")
    train_lines = [
        "накопление примеров жеста",
        "сохранение обучающей выборки",
        "обновление модели пользователя",
    ]
    for i, line in enumerate(train_lines):
        draw.text((290, train_y + 82 + i * 30), line, font=small_font, fill="#455A64")

    infer_x = 980
    draw.rounded_rectangle((infer_x, train_y, 1570, train_y + 170), radius=18, fill="#F7F9FB", outline="#78909C", width=2)
    draw.text((infer_x + 40, train_y + 30), "Режим распознавания", font=box_font, fill="#1F2933")
    infer_lines = [
        "предсказание класса жеста",
        "сглаживание результата",
        "запуск привязанного действия",
    ]
    for i, line in enumerate(infer_lines):
        draw.text((infer_x + 60, train_y + 82 + i * 30), line, font=small_font, fill="#455A64")

    arrow(draw, centers[3], (520, train_y - 10))
    arrow(draw, centers[4], (1275, train_y - 10))
    img.save(path)


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, before=0, after=0, line=1.15):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(1.25)
    else:
        fmt.first_line_indent = Cm(0)


def add_text_paragraph(doc, text, first_line=True):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=first_line)
    run = p.add_run(text)
    set_run_font(run, 14)
    return p


def add_center(doc, text, size=14, bold=False, after=0):
    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=2, after=6, line=1.0)
    run = p.add_run(text)
    set_run_font(run, 12, italic=True)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_image(doc, image_path, width_cm=16.5):
    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=6, after=0, line=1.0)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.font.size = Pt(14)

    for name in ["List Number", "List Bullet"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0


def build_docx():
    ASSET_DIR.mkdir(exist_ok=True)
    arch = ASSET_DIR / "architecture.png"
    flow = ASSET_DIR / "data_flow.png"
    make_architecture_image(arch)
    make_flow_image(flow)

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    add_page_field(header_p)

    add_center(doc, "Р. Т. Давар", size=14)
    add_center(doc, "студент кафедры прикладной информатики", size=14)
    add_center(doc, "Т. М. Татарникова – д-р техн. наук, профессор – научный руководитель", size=14, after=4)
    add_center(doc, "Система бесконтактного взаимодействия человека с компьютером на основе методов компьютерного зрения", size=14, bold=True, after=6)

    paragraphs = [
        "Развитие интеллектуальных пользовательских интерфейсов связано с поиском способов взаимодействия, которые не требуют постоянного использования клавиатуры, мыши или сенсорного экрана. В ряде сценариев традиционные устройства ввода оказываются недостаточно удобными: пользователь может работать с мультимедийной системой на расстоянии, управлять приложением во время демонстрации, выполнять команды в условиях ограниченного доступа к периферии или использовать компьютер в режиме повышенной стерильности. В таких случаях актуальной становится задача бесконтактного управления на основе жестов руки.",
        "Целью работы является разработка системы распознавания жестов руки в реальном времени на основе методов компьютерного зрения. Объектом исследования выступает процесс бесконтактного взаимодействия человека с компьютером с использованием жестов руки. Предметом исследования являются архитектура и программные решения системы, обеспечивающие выделение признаков жеста, обучение пользовательской модели и последующее распознавание в режиме реального времени.",
        "Существующие подходы к распознаванию жестов условно можно разделить на две группы: wearable/glove-based и vision-based. Первый подход предполагает применение специальных перчаток, датчиков или иных носимых устройств, что может повышать точность измерений, но снижает естественность взаимодействия и усложняет эксплуатацию. В данной работе выбран vision-based подход, так как он не требует дополнительного оборудования, кроме обычной камеры, лучше соответствует идее естественного интерфейса и проще интегрируется в прикладные компьютерные системы.",
        "В рамках прототипа основное внимание уделено статическим и квазистатическим жестам. Такой выбор позволяет упростить этап обучения, уменьшить требования к объему пользовательской выборки и обеспечить стабильное распознавание в условиях реального времени. При этом архитектура системы сохраняет возможность дальнейшего расширения: при необходимости в нее могут быть добавлены динамические признаки, анализ временных последовательностей и более сложные модели классификации.",
    ]
    for text in paragraphs:
        add_text_paragraph(doc, text)

    add_image(doc, arch)
    add_caption(doc, "Рис. 1. Общая архитектура системы распознавания жестов")

    more = [
        "Общая архитектура системы представлена на рис. 1. На первом этапе выполняется захват видеопотока с камеры. Затем модуль компьютерного зрения выделяет ключевые точки кисти, которые описывают положение основных суставов руки. Полученные координаты преобразуются в компактное признаковое представление: выполняется нормализация относительно положения кисти, после чего формируется числовой вектор, пригодный для передачи в классификатор.",
        "Отдельным элементом системы является словарь пользовательских жестов. Пользователь может добавлять новые жесты без изменения программного кода: система сохраняет обучающие примеры, связывает их с заданным названием и использует при последующем обучении модели. Такой подход делает приложение персонализируемым, поскольку жестовый словарь может быть адаптирован под конкретного пользователя и конкретные команды.",
        "Принципиальным архитектурным решением стало разделение режимов работы. Обучение модели и распознавание жестов выполняются как логически самостоятельные процессы. В режиме обучения система накапливает примеры, формирует обучающую выборку и обновляет классификатор. В режиме распознавания видеопоток обрабатывается непрерывно, а результат классификации используется для запуска привязанного действия. Благодаря этому пользовательский интерфейс не содержит логики компьютерного зрения и машинного обучения, а вычислительные модули могут развиваться независимо от визуальной части приложения.",
    ]
    for text in more:
        add_text_paragraph(doc, text)

    add_image(doc, flow)
    add_caption(doc, "Рис. 2. Поток данных при обучении и распознавании жестов")

    final_paragraphs = [
        "Поток данных в системе показан на рис. 2. Каждый кадр с камеры проходит этап выделения кисти и ключевых точек. Далее координаты переводятся в единое представление, устойчивое к смещению руки в кадре и изменению масштаба. После этого данные передаются в один из режимов работы: сохранение нового жеста, обучение модели или распознавание уже обученного жеста.",
        "В программной реализации прототипа были реализованы функции захвата видеопотока, выделения ключевых точек руки, добавления пользовательских жестов, накопления обучающих примеров, обучения модели и распознавания жестов в реальном времени. Типовой сценарий работы пользователя включает запуск приложения, активацию камеры, добавление нового жеста, запись нескольких примеров, обучение модели и переход к распознаванию. При распознавании жест может быть связан с прикладной командой, что превращает классификатор из демонстрационного модуля в инструмент управления компьютерной системой.",
        "Ключевым преимуществом предложенной системы является модульность. Модуль захвата видео отвечает только за получение кадров, модуль компьютерного зрения — за выделение информативных точек кисти, модуль признаков — за преобразование координат, модуль машинного обучения — за классификацию, а интерфейс — за пользовательский сценарий. Такое разделение снижает связанность компонентов и упрощает расширение: можно заменить классификатор, добавить новые команды или изменить интерфейс без полной переработки архитектуры.",
        "В результате работы был разработан работоспособный прототип системы бесконтактного взаимодействия человека с компьютером. Реализован полный цикл работы с жестами: от добавления пользовательского жеста до его распознавания в реальном времени. Полученные результаты подтверждают применимость модульной архитектуры для создания прикладного жестового интерфейса, ориентированного на обычную веб-камеру и не требующего специализированных носимых устройств.",
        "Вместе с тем система имеет ряд ограничений. Качество распознавания зависит от освещения, контрастности фона, положения руки относительно камеры и стабильности выделения ключевых точек. Кроме того, для устойчивой работы требуется достаточный объем обучающих примеров, а текущая версия ориентирована преимущественно на статические и квазистатические жесты. Дальнейшее развитие работы может быть связано с расширением обучающей выборки, поддержкой динамических жестов, сравнением нескольких классификаторов и внедрением адаптивной оптимизации параметров модели под конкретного пользователя.",
        "Таким образом, использование методов компьютерного зрения позволяет построить удобный и расширяемый бесконтактный интерфейс управления компьютером. Предложенная система решает задачу распознавания пользовательских жестов в реальном времени и может применяться как основа для персонализируемых средств управления прикладными программами.",
    ]
    for text in final_paragraphs:
        add_text_paragraph(doc, text)

    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=8, after=4)
    r = p.add_run("Библиографический список")
    set_run_font(r, size=14, bold=True)

    sources = [
        "MediaPipe Hand Landmarker task guide [Электронный ресурс]. – URL: https://ai.google.dev/edge/mediapipe/solutions/vision/hand_landmarker (дата обращения: 30.05.2026).",
        "OpenCV documentation. VideoCapture class [Электронный ресурс]. – URL: https://docs.opencv.org/4.x/d8/dfe/classcv_1_1VideoCapture.html (дата обращения: 30.05.2026).",
        "Scikit-learn documentation. KNeighborsClassifier [Электронный ресурс]. – URL: https://scikit-learn.org/stable/modules/generated/sklearn.neighbors.KNeighborsClassifier.html (дата обращения: 30.05.2026).",
        "Optuna documentation. A hyperparameter optimization framework [Электронный ресурс]. – URL: https://optuna.readthedocs.io/ (дата обращения: 30.05.2026).",
    ]
    for source in sources:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(source)
        set_run_font(run, size=12)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_docx())
