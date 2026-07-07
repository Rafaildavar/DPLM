from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT / "docs" / "Диплом все - с графиками Tableau.docx"


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def insert_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def find_exact(doc: Document, text: str, start: int = 0) -> tuple[int, Paragraph]:
    for index, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip() == text:
            return index, paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_startswith(doc: Document, prefix: str, start: int = 0) -> tuple[int, Paragraph]:
    for index, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip().startswith(prefix):
            return index, paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def set_title(paragraph: Paragraph, text: str) -> None:
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = True


def format_body(paragraph: Paragraph, first_line: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.49)


def replace_between(
    doc: Document,
    start_text: str,
    end_text: str,
    new_title: str,
    body: list[str],
    start_offset: int = 0,
    page_break_after: bool = True,
) -> None:
    start_index, start = find_exact(doc, start_text, start=start_offset)
    end_index, end = find_exact(doc, end_text, start=start_index + 1)
    for paragraph in list(doc.paragraphs[start_index + 1 : end_index]):
        delete_paragraph(paragraph)

    set_title(start, new_title)
    last = start
    for text in body:
        last = insert_after(last, text)
        format_body(last, first_line=not text.startswith(("Ключевые слова:", "Объект", "Предмет", "Цель", "Методы", "Практическая")))
    if page_break_after:
        last.add_run().add_break(WD_BREAK.PAGE)


def replace_intro(doc: Document, body: list[str]) -> None:
    start_index, start = find_exact(doc, "ВВЕДЕНИЕ", start=90)
    end_index, end = find_startswith(doc, "1 Теоретические сведения", start=start_index + 1)
    for paragraph in list(doc.paragraphs[start_index + 1 : end_index]):
        delete_paragraph(paragraph)

    set_title(start, "ВВЕДЕНИЕ")
    last = start
    for text in body:
        last = insert_after(last, text)
        format_body(last)
    last.add_run().add_break(WD_BREAK.PAGE)


def replace_references(doc: Document, references: list[str]) -> None:
    start_index, start = find_exact(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for paragraph in list(doc.paragraphs[start_index + 1 :]):
        delete_paragraph(paragraph)
    set_title(start, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    last = start
    for item in references:
        last = insert_after(last, item)
        last.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        last.paragraph_format.left_indent = Inches(0.25)
        last.paragraph_format.first_line_indent = Inches(-0.25)


def replace_para(doc: Document, prefix: str, text: str) -> None:
    _, paragraph = find_startswith(doc, prefix)
    paragraph.text = text
    format_body(paragraph)


def update_toc(doc: Document) -> None:
    page_numbers = {
        "ВВЕДЕНИЕ": 7,
        "1 Теоретические сведения и анализ предметной области": 9,
        "1.1 Человеко-компьютерное взаимодействие и жестовое управление": 9,
        "1.2 Современные подходы к распознаванию жестов": 10,
        "1.3 Методы детекции руки и извлечения ключевых точек": 12,
        "1.4 Методы классификации жестов и обучение на пользовательских данных": 15,
        "1.5 Обзор существующих решений и требований к дружелюбному интерфейсу": 18,
        "1.6 Выводы по первой главе": 22,
        "2 Анализ требований и проектирование системы": 24,
        "2.1 Постановка задачи и сценарии использования": 24,
        "2.2 Функциональные и нефункциональные требования": 25,
        "2.3 Выбор технологического стека": 27,
        "2.4 Общая архитектура системы GestureBind": 28,
        "2.5 Проектирование информационной модели и базы данных": 30,
        "2.6 Проектирование привязки жестов к командам": 31,
        "2.7 Проектирование пользовательского интерфейса и диагностики": 33,
        "2.8 Выводы по второй главе": 34,
        "3 Разработка и реализация системы GestureBind": 35,
        "3.1 Организация разработки и структура программных модулей": 36,
        "3.2 Реализация захвата видеопотока и выделения ключевых точек руки": 36,
        "3.3 Нормализация landmarks и формирование признакового вектора": 38,
        "3.4 Реализация записи пользовательских обучающих примеров": 39,
        "3.5 Реализация обучения классификатора жестов": 41,
        "3.6 Реализация онлайн-распознавания и стабилизации результата": 42,
        "3.7 Реализация выполнения команд по распознанным жестам": 44,
        "3.8 Реализация пользовательского интерфейса и прикладного контроллера": 45,
        "3.9 Выводы по третьей главе": 46,
        "4 Разработка программного приложения и пользовательского интерфейса": 48,
        "4.1 Назначение прикладного интерфейса": 48,
        "4.2 Выбор интерфейсного слоя и организация окна приложения": 49,
        "4.3 Прикладной контроллер и событийная модель": 50,
        "4.4 Главный экран распознавания": 52,
        "4.5 Экран обучения и экран пользовательского словаря": 53,
        "4.6 Экран привязки жестов к командам": 54,
        "4.7 Настройки, конфигурация и диагностика": 55,
        "4.8 Хранение состояния приложения и журналирование": 57,
        "4.9 Выводы по четвертой главе": 58,
        "5 Тестирование и оценка результатов": 59,
        "5.1 Методика тестирования": 59,
        "5.2 Автоматизированное модульное тестирование": 60,
        "5.3 Проверка пользовательских сценариев": 61,
        "5.4 Оценка качества распознавания жестов": 63,
        "5.5 Оценка производительности и оптимизации": 64,
        "5.6 Ограничения разработанного решения": 66,
        "5.7 Направления дальнейшего развития": 67,
        "5.8 Выводы по пятой главе": 68,
        "ЗАКЛЮЧЕНИЕ": 70,
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ": 73,
    }
    content_index, _ = find_exact(doc, "СОДЕРЖАНИЕ")
    terms_index, _ = find_exact(doc, "УСЛОВНЫЕ ОБОЗНАЧЕНИЯ")
    for paragraph in doc.paragraphs[content_index + 1 : terms_index]:
        text = paragraph.text.strip()
        if not text or "\t" not in paragraph.text:
            continue
        title = paragraph.text.rsplit("\t", 1)[0]
        if title in page_numbers:
            paragraph.text = f"{title}\t{page_numbers[title]}"


REFERAT = [
    "Пояснительная записка к выпускной квалификационной работе содержит {{PAGES}} с., 34 табл., 21 рис., 42 источника.",
    "Ключевые слова: жестовое управление, компьютерное зрение, MediaPipe Hand Landmarker, KNN, пользовательский словарь жестов, Flet, ORM, Tableau Public.",
    "Объект исследования - методы и программные средства человеко-компьютерного взаимодействия, основанные на компьютерном зрении и машинном обучении.",
    "Предмет исследования - desktop-система распознавания пользовательских жестов с механизмом записи собственного словаря, обучения классификатора и привязки жестов к командам компьютера.",
    "Цель работы - разработать систему GestureBind для записи пользовательских жестов, обучения классификатора, распознавания жестов в реальном времени, выполнения связанных команд и отображения диагностической обратной связи.",
    "Методы работы включают анализ предметной области, проектирование архитектуры, извлечение landmarks руки, классификацию признаков методом k-ближайших соседей, ORM-моделирование данных, модульное тестирование и визуальный анализ результатов в Tableau Public.",
    "В результате разработан прототип приложения GestureBind. Реализованы запись обучающих примеров, обучение KNN-модели, онлайн-распознавание жестов, хранение жестов, команд и сэмплов в БД, интерфейс обучения и привязки команд.",
    "Практическая значимость состоит в возможности использовать приложение как основу персонализируемого жестового интерфейса для управления компьютером, презентациями, мультимедиа и дальнейших исследований пользовательских жестовых словарей.",
]


TERMS = [
    "1. CV - Computer Vision, компьютерное зрение;",
    "2. ML - Machine Learning, машинное обучение;",
    "3. HCI - Human-Computer Interaction, человеко-компьютерное взаимодействие;",
    "4. GUI - Graphical User Interface, графический пользовательский интерфейс;",
    "5. UI - User Interface, пользовательский интерфейс;",
    "6. UX - User Experience, пользовательский опыт;",
    "7. KNN - k-Nearest Neighbors, метод k-ближайших соседей;",
    "8. SVM - Support Vector Machine, метод опорных векторов;",
    "9. RGB - Red, Green, Blue, цветовая модель изображения;",
    "10. FPS - Frames Per Second, количество кадров в секунду;",
    "11. ROI - Region of Interest, область интереса на изображении;",
    "12. Landmark - ключевая точка руки, возвращаемая моделью MediaPipe;",
    "13. Handedness - признак принадлежности обнаруженной руки к левой или правой;",
    "14. Feature vector - признаковый вектор, подаваемый на вход классификатору;",
    "15. Inference - инференс, применение обученной модели к новым данным;",
    "16. Confidence - численная оценка уверенности модели в результате распознавания;",
    "17. Cooldown - временная задержка между повторными выполнениями команды;",
    "18. ORM - Object-Relational Mapping, объектно-реляционное отображение;",
    "19. CRUD - Create, Read, Update, Delete, базовые операции с данными;",
    "20. БД - база данных;",
    "21. API - Application Programming Interface, программный интерфейс приложения;",
    "22. CLI - Command Line Interface, интерфейс командной строки;",
    "23. JSON - JavaScript Object Notation, формат структурированного обмена данными;",
    "24. NumPy - библиотека Python для работы с массивами и численными данными;",
    "25. Pytest - фреймворк автоматизированного тестирования Python;",
    "26. BI - Business Intelligence, инструменты аналитической визуализации данных;",
    "27. Tableau Public - облачный сервис Tableau для построения и публикации визуализаций.",
]


INTRO = [
    "Современные программные системы все чаще используют естественные способы взаимодействия с пользователем. Помимо клавиатуры и мыши применяются сенсорное управление, распознавание позы, голоса и жестов. Такие способы ввода позволяют сделать работу с компьютером более гибкой и доступной, особенно в сценариях презентаций, мультимедиа, hands-free управления, ассистивных интерфейсов и демонстрационных систем.",
    "Одним из перспективных направлений является распознавание жестов руки с использованием компьютерного зрения. При наличии обычной веб-камеры система может получать видеопоток, выделять ключевые точки руки, классифицировать жест и выполнять связанную с ним команду. Теоретическую основу такого подхода составляют методы обработки изображений и компьютерного зрения, описанные в классических работах по Computer Vision и в библиотечных реализациях OpenCV [25, 26].",
    "Актуальность темы определяется тем, что большинство простых жестовых решений ориентированы на заранее заданный набор жестов. Такой подход удобен для демонстрации алгоритма, но хуже подходит для реального пользователя: жесты могут отличаться по привычкам, моторике, условиям съемки и назначаемым действиям. Поэтому в данной работе рассматривается система, в которой пользователь может сформировать собственный словарь, записать обучающие примеры и самостоятельно связать распознанный жест с нужной командой.",
    "Важной особенностью выбранной задачи является сочетание нескольких направлений разработки. Система должна не только распознавать руку в кадре, но и хранить пользовательские данные, синхронизировать файлы обучающих примеров с базой данных, обучать классификатор на малом наборе, показывать состояние камеры и модели, а также безопасно выполнять команды операционной системы. Поэтому проект GestureBind рассматривается как прикладное desktop-приложение, а не только как изолированный эксперимент по классификации.",
    "Целью выпускной квалификационной работы является разработка системы с функцией обучения для распознавания жестов на основе собственного словаря пользователя. Для достижения цели необходимо проанализировать предметную область, выбрать технологический стек, спроектировать архитектуру приложения, реализовать модуль распознавания жестов, разработать интерфейс записи и обучения, обеспечить привязку жестов к командам, организовать хранение данных и провести тестирование разработанного решения.",
    "Объектом исследования являются методы и программные средства человеко-компьютерного взаимодействия, основанные на компьютерном зрении и машинном обучении. Предметом исследования является разработка desktop-системы распознавания пользовательских жестов с механизмом обучения собственного словаря, привязки команд и диагностической обратной связи.",
    "В качестве базовых технологий в работе используются Python, OpenCV, MediaPipe Hand Landmarker, NumPy, scikit-learn, Flet, SQLAlchemy ORM, PostgreSQL/SQLite, PyAutoGUI и pytest. MediaPipe применяется для получения landmarks руки, scikit-learn - для обучения KNN-классификатора, Flet - для реализации графического интерфейса, SQLAlchemy ORM - для работы с сущностями приложения, а pytest - для автоматизированной проверки модулей. Выбор этих инструментов соответствует практической задаче создания локального приложения, которое можно запустить на обычном компьютере без специализированного оборудования [27-39].",
    "Методика работы включает анализ научных и технических источников, проектирование требований, разработку архитектуры, реализацию программных модулей, запись пользовательских сэмплов, обучение классификатора, проверку пользовательских сценариев и анализ качества распознавания. Для усиления раздела тестирования результаты метрик и распределения проверок дополнительно визуализируются в Tableau Public, что позволяет представить оценку системы в виде понятных BI-графиков.",
    "Практическая значимость работы состоит в том, что разработанный прототип может использоваться как основа персонализируемого жестового интерфейса. Пользователь не ограничивается фиксированным набором действий: он записывает собственные жесты, обучает модель и назначает команды под свои задачи. Такой подход соответствует принципам человеко-компьютерного взаимодействия, где важны понятность сценария, обратная связь, предотвращение ошибок и контроль пользователя над поведением системы [40-42].",
    "Выпускная квалификационная работа состоит из введения, пяти глав, заключения и списка использованных источников. В первой главе рассматриваются теоретические сведения и существующие подходы к распознаванию жестов. Во второй главе формулируются требования и проектируется архитектура GestureBind. В третьей главе описывается реализация ядра системы. В четвертой главе рассматривается пользовательский интерфейс. В пятой главе приводятся результаты тестирования, метрики распознавания, визуализации Tableau Public и ограничения разработанного решения.",
]


REFERENCES = [
    "1. Oudah M., Al-Naji A., Chahl J. Hand Gesture Recognition Based on Computer Vision: A Review of Techniques. Journal of Imaging, 2020. URL: https://www.mdpi.com/2313-433X/6/8/73",
    "2. Linardakis M., Varlamis I., Papadopoulos G. Th. Survey on Hand Gesture Recognition from Visual Input. arXiv:2501.11992, 2025. URL: https://arxiv.org/abs/2501.11992",
    "3. Zhang F., Bazarevsky V., Vakunov A., Tkachenka A., Sung G., Chang C.-L., Grundmann M. MediaPipe Hands: On-device Real-time Hand Tracking. arXiv:2006.10214, 2020. URL: https://arxiv.org/abs/2006.10214",
    "4. Google AI for Developers. Hand landmarks detection guide. URL: https://ai.google.dev/edge/mediapipe/solutions/vision/hand_landmarker",
    "5. scikit-learn. Nearest Neighbors documentation. URL: https://scikit-learn.org/stable/modules/neighbors.html",
    "6. Akiba T., Sano S., Yanase T., Ohta T., Koyama M. Optuna: A Next-generation Hyperparameter Optimization Framework. arXiv:1907.10902, 2019. URL: https://arxiv.org/abs/1907.10902",
    "7. Bischl B., Binder M., Lang M. et al. Hyperparameter Optimization: Foundations, Algorithms, Best Practices and Open Challenges. arXiv:2107.05847, 2021. URL: https://arxiv.org/abs/2107.05847",
    "8. Nielsen J. 10 Usability Heuristics for User Interface Design. Nielsen Norman Group, updated 2024. URL: https://www.nngroup.com/articles/ten-usability-heuristics/",
    "9. Безед К. Методика построения человеко-машинного интерфейса для управления компьютером на основе библиотеки MediaPipe // Молодой исследователь Дона. 2023. No. 1(40). С. 5-10.",
    "10. Moysiadis V. et al. An Integrated Real-Time Hand Gesture Recognition Framework for Human-Robot Interaction in Agriculture // Applied Sciences. 2022. Vol. 12, No. 16. Article 8160. DOI: 10.3390/app12168160.",
    "11. Gao K. et al. Challenges and solutions for vision-based hand gesture interpretation: A review // Computer Vision and Image Understanding. 2024. Vol. 248. Article 104095. DOI: 10.1016/j.cviu.2024.104095.",
    "12. Kapitanov A., Kvanchiani K., Nagaev A., Kraynov R., Makhliarchuk A. HaGRID - HAnd Gesture Recognition Image Dataset // 2024 IEEE/CVF Winter Conference on Applications of Computer Vision (WACV). 2024. P. 4560-4569. DOI: 10.1109/WACV57701.2024.00451.",
    "13. Fertl E., Castillo E., Stettinger G., Cuellar M. P., Morales D. P. Hand Gesture Recognition on Edge Devices: Sensor Technologies, Algorithms, and Processing Hardware // Sensors. 2025. Vol. 25, No. 6. Article 1687. DOI: 10.3390/s25061687.",
    "14. Popov P. A., Laganiere R. Long Hands gesture recognition system: 2 step gesture recognition with machine learning and geometric shape analysis // Multimedia Tools and Applications. 2022. Vol. 81. P. 40311-40342. DOI: 10.1007/s11042-022-12870-8.",
    "15. Uboweja E. et al. On-device Real-time Custom Hand Gesture Recognition. arXiv:2309.10858, 2023. DOI: 10.48550/arXiv.2309.10858.",
    "16. Sung G. et al. On-device Real-time Hand Gesture Recognition. arXiv:2111.00038, 2021. DOI: 10.48550/arXiv.2111.00038.",
    "17. Meng Y., Jiang H., Duan N., Wen H. Real-Time Hand Gesture Monitoring Model Based on MediaPipe's Registerable System // Sensors. 2024. Vol. 24, No. 19. Article 6262. DOI: 10.3390/s24196262.",
    "18. Uddin M. Z., Boletsis C., Rudshavn P. Real-Time Norwegian Sign Language Recognition Using MediaPipe and LSTM // Multimodal Technologies and Interaction. 2025. Vol. 9, No. 3. Article 23. DOI: 10.3390/mti9030023.",
    "19. ZainEldin H. et al. Silent no more: a comprehensive review of artificial intelligence, deep learning, and machine learning in facilitating deaf and mute communication // Artificial Intelligence Review. 2024. Vol. 57, No. 7. Article 188. DOI: 10.1007/s10462-024-10816-0.",
    "20. Gil-Martin M., Marini M., Martin-Fernandez I., Esteban-Romero S., Cinque L. Hand Gesture Recognition Using MediaPipe Landmarks and Deep Learning Networks // Proceedings of the 17th International Conference on Agents and Artificial Intelligence. 2025. P. 24-30. DOI: 10.5220/0013053500003890.",
    "21. Apple Inc. Method and device for defining custom hand gestures. Patent US12045392B1. 2024.",
    "22. Microsoft Technology Licensing LLC. Architecture for controlling a computer using hand gestures. Patent US9652042B2. 2017.",
    "23. PointGrab Ltd. Computer vision gesture based control of a device. Patent WO2011045789A1. 2011.",
    "24. Intel Corporation. Hand gesture recognition system. Patent US8781221B2. 2014.",
    "25. Szeliski R. Computer Vision: Algorithms and Applications. 2nd ed. Springer, 2022. URL: https://szeliski.org/Book/ (дата обращения: 23.05.2026).",
    "26. Bradski G. The OpenCV Library // Dr. Dobb's Journal of Software Tools. 2000. Vol. 25, No. 11. P. 120-125.",
    "27. Lugaresi C., Tang J., Nash H. et al. MediaPipe: A Framework for Building Perception Pipelines [Электронный ресурс]. URL: https://arxiv.org/abs/1906.08172 (дата обращения: 23.05.2026).",
    "28. Google AI Edge. Hand landmarks detection guide [Электронный ресурс]. URL: https://ai.google.dev/edge/mediapipe/solutions/vision/hand_landmarker (дата обращения: 23.05.2026).",
    "29. OpenCV Documentation [Электронный ресурс]. URL: https://docs.opencv.org/ (дата обращения: 20.05.2026).",
    "30. scikit-learn. Nearest Neighbors [Электронный ресурс]. URL: https://scikit-learn.org/stable/modules/neighbors.html (дата обращения: 23.05.2026).",
    "31. scikit-learn. KNeighborsClassifier [Электронный ресурс]. URL: https://scikit-learn.org/stable/modules/generated/sklearn.neighbors.KNeighborsClassifier.html (дата обращения: 23.05.2026).",
    "32. NumPy Documentation [Электронный ресурс]. URL: https://numpy.org/doc/ (дата обращения: 20.05.2026).",
    "33. Python 3 Documentation [Электронный ресурс]. URL: https://docs.python.org/3/ (дата обращения: 20.05.2026).",
    "34. Flet Documentation [Электронный ресурс]. URL: https://flet.dev/docs/ (дата обращения: 25.05.2026).",
    "35. SQLAlchemy ORM Documentation [Электронный ресурс]. URL: https://docs.sqlalchemy.org/20/orm/ (дата обращения: 13.05.2026).",
    "36. PostgreSQL Documentation [Электронный ресурс]. URL: https://www.postgresql.org/docs/ (дата обращения: 13.05.2026).",
    "37. Alembic Documentation [Электронный ресурс]. URL: https://alembic.sqlalchemy.org/ (дата обращения: 13.05.2026).",
    "38. PyAutoGUI Documentation [Электронный ресурс]. URL: https://pyautogui.readthedocs.io/ (дата обращения: 20.05.2026).",
    "39. pytest Documentation [Электронный ресурс]. URL: https://docs.pytest.org/en/stable/ (дата обращения: 20.05.2026).",
    "40. Dix A., Finlay J., Abowd G., Beale R. Human-Computer Interaction. 3rd ed. Pearson, 2004.",
    "41. Nielsen J. Usability Engineering. San Francisco: Morgan Kaufmann, 1993.",
    "42. Fowler M. Patterns of Enterprise Application Architecture. Boston: Addison-Wesley, 2002.",
]


def main() -> None:
    doc = Document(DOC_PATH)

    replace_between(doc, "АННОТАЦИЯ", "СОДЕРЖАНИЕ", "РЕФЕРАТ", REFERAT)
    replace_between(doc, "УСЛОВНЫЕ ОБОЗНАЧЕНИЯ", "ВВЕДЕНИЕ", "УСЛОВНЫЕ ОБОЗНАЧЕНИЯ", TERMS, start_offset=80)
    replace_intro(doc, INTRO)

    replace_para(
        doc,
        "Жестовое управление особенно актуально",
        "Жестовое управление особенно актуально в сценариях, где использование традиционных устройств ввода затруднено или менее удобно. К таким сценариям относятся проведение презентаций, управление мультимедиа, работа в режиме hands-free, ассистивные интерфейсы, управление умным домом, взаимодействие с AR/VR-средами и робототехническими системами. Обзорные исследования по распознаванию жестов отмечают, что жесты применяются в human-computer interaction, робототехнике, домашней автоматизации, медицинских и коммуникационных приложениях [1, 40, 41].",
    )
    replace_para(
        doc,
        "Распознавание жестов руки может решаться",
        "Распознавание жестов руки может решаться различными методами в зависимости от того, какие данные доступны системе и какие ограничения накладываются на пользователя. В обзорной литературе обычно выделяют два крупных класса подходов: сенсорные и визуальные [1]. Сенсорные подходы предполагают использование перчаток, браслетов, инерциальных датчиков или других устройств, которые непосредственно измеряют положение кисти и пальцев. Визуальные подходы используют камеру и алгоритмы компьютерного зрения, опирающиеся на методы обработки изображений и выделения признаков [25, 26].",
    )
    replace_para(
        doc,
        "В GestureBind для этой задачи используется MediaPipe Hands.",
        "В GestureBind для этой задачи используется MediaPipe Hands. MediaPipe как фреймворк для построения perception pipelines описывает модульный подход к сборке конвейеров компьютерного зрения и обработки медиаданных [27]. В статье \"MediaPipe Hands: On-device Real-time Hand Tracking\" описан конвейер, который строит скелет руки по изображению с одной RGB-камеры. Конвейер состоит из двух моделей: детектора ладони и модели ключевых точек руки [3]. Такой подход позволяет сначала локализовать область руки, а затем выполнять более точную регрессию координат внутри найденной области.",
    )
    replace_para(
        doc,
        "Официальная документация MediaPipe Hand Landmarker",
        "Официальная документация MediaPipe Hand Landmarker указывает, что задача позволяет обнаруживать ключевые точки рук на изображении, работать со статическими изображениями, видеокадрами и live-видеопотоком, а на выходе возвращает координаты landmarks, мировые координаты и информацию о том, левая или правая рука обнаружена [4, 28]. Для приложений реального времени важно, что palm detector не обязан запускаться на каждом кадре. При успешном отслеживании используется информация из предыдущего кадра, а повторная детекция выполняется только при потере руки.",
    )
    replace_para(
        doc,
        "Метод k-ближайших соседей",
        "Метод k-ближайших соседей (KNN) является одним из наиболее понятных вариантов для классификации пользовательских жестов. Согласно документации scikit-learn, neighbors-based classification относится к instance-based learning: модель не строит сложного внутреннего обобщения, а хранит обучающие примеры и классифицирует новый объект по большинству классов среди ближайших соседей [5, 30, 31]. Для GestureBind это удобно по нескольким причинам:",
    )
    replace_para(
        doc,
        "Отдельной частью задачи является оптимизация гиперпараметров.",
        "Отдельной частью задачи является проверка влияния параметров и предварительной обработки признаков на качество классификации. Ручной подбор параметров может быть трудоемким и плохо воспроизводимым. Обзор по hyperparameter optimization отмечает, что многие алгоритмы машинного обучения имеют гиперпараметры, существенно влияющие на результат, а автоматическая оптимизация позволяет уйти от ручного trial-and-error подхода [7]. В текущей реализации GestureBind основной воспроизводимый эксперимент выполнен для нескольких вариантов обработки признаков KNN, а полноценный автоматический подбор параметров рассматривается как направление дальнейшего развития.",
    )
    replace_para(
        doc,
        "В качестве языка реализации используется Python.",
        "В качестве языка реализации выбран Python, поскольку он имеет развитую экосистему библиотек для компьютерного зрения, машинного обучения, работы с массивами, GUI, ORM и тестирования [33]. Для захвата кадров и базовой обработки изображений применяется OpenCV [29], для численных массивов и сохранения обучающих примеров - NumPy [32], для классификации жестов - scikit-learn [30, 31]. Пользовательский интерфейс реализуется на Flet [34], слой хранения данных - на SQLAlchemy ORM с возможностью использования PostgreSQL и миграций Alembic [35, 36, 37], выполнение пользовательских действий опирается на PyAutoGUI [38], а автоматизированные проверки выполняются с помощью pytest [39].",
    )
    replace_para(
        doc,
        "Файловая система используется для данных",
        "Файловая система используется для данных, которые удобно читать и записывать как массивы или артефакты модели. К ним относятся NumPy-файлы обучающих примеров, файл обученной модели, список классов и размерность признаков. База данных используется для структурированных сущностей: жестов, сэмплов, моделей, команд, настроек и журналов распознавания. Такое разделение соответствует практическому подходу к построению прикладной архитектуры, где файловые артефакты и доменные сущности обслуживаются разными механизмами хранения [35, 42].",
    )
    replace_para(
        doc,
        "В пятой главе была проведена оценка результатов разработки GestureBind.",
        "В пятой главе была проведена оценка результатов разработки GestureBind. Описана методика тестирования, включающая модульные проверки, пользовательские сценарии, оценку качества классификации и анализ ограничений.",
    )
    replace_para(
        doc,
        "Стабильный набор автоматизированных unit-тестов по ядру жестовой системы показал",
        "Стабильный набор автоматизированных unit-тестов по ядру жестовой системы показал результат 137 из 137 пройденных тестов. Проверки охватывают правила привязок и безопасности, команды и системные действия, хранение данных, CV-модули, Flet-контроллер, UI-тексты, управление указателем и голосовой ассистент.",
    )
    replace_para(
        doc,
        "Оценка распознавания на доступной части локального набора данных показала",
        "Оценка распознавания на текущем локальном наборе данных из 41 примера и двух классов показала, что базовый pipeline KNN дает accuracy 0.9268 и macro F1 0.9261, а вариант с центрированием признаков каждого примера достигает 1.0000 по основным метрикам. Результат подтверждает работоспособность реализованного pipeline, но требует осторожной интерпретации из-за малого числа классов и ограниченного объема выборки.",
    )
    replace_para(
        doc,
        "В пятой главе была выполнена оценка результатов разработки.",
        "В пятой главе была выполнена оценка результатов разработки. Стабильный набор unit-тестов по ядру жестовой системы показал результат 137 из 137 пройденных тестов. Оценка классификации на текущем локальном наборе данных показала accuracy 0.9268 и macro F1 0.9261 для базового pipeline KNN, а лучший проверенный вариант предварительной обработки дал 1.0000 по основным метрикам. Полученные значения подтверждают работоспособность реализованного решения, однако требуют осторожной интерпретации из-за малого числа классов и ограниченного объема выборки.",
    )

    replace_references(doc, REFERENCES)
    update_toc(doc)

    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
