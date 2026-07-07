from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CHAPTER_MDS = [
    ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CHAPTER_1.md",
    ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CHAPTER_2.md",
    ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CHAPTER_3.md",
    ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CHAPTER_4.md",
    ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CHAPTER_5.md",
]
CONCLUSION_MD = ROOT / "docs" / "diploma" / "drafts" / "DIPLOMA_REPORT_CONCLUSION.md"
APPENDICES = [
    {
        "label": "А",
        "title": "Фрагмент загрузки датасета и обучения классификатора",
        "source": ROOT / "cv" / "train_classifier.py",
        "start": 17,
        "end": 98,
    },
    {
        "label": "Б",
        "title": "Фрагмент онлайн-распознавания жестов",
        "source": ROOT / "app" / "gesture_online_infer.py",
        "start": 264,
        "end": 372,
    },
    {
        "label": "В",
        "title": "Фрагмент политики привязки жеста к команде",
        "source": ROOT / "app" / "services" / "gesture_command_bridge.py",
        "start": 68,
        "end": 182,
    },
    {
        "label": "Г",
        "title": "Фрагмент интерфейса привязки жестов к командам",
        "source": ROOT / "app" / "flet_app" / "views" / "bindings.py",
        "start": 147,
        "end": 275,
    },
]
OUT_DOCX = ROOT / "docs" / "Диплом_Давар_GestureBind_черновик.docx"


def set_run_font(run, size: float = 14, bold: bool = False, italic: bool = False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Times New Roman")


def set_run_code_font(run, size: float = 8):
    run.font.name = "Courier New"
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Courier New")
    r_fonts.set(qn("w:hAnsi"), "Courier New")
    r_fonts.set(qn("w:cs"), "Courier New")
    r_fonts.set(qn("w:eastAsia"), "Courier New")


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_width(table, widths_cm: list[float]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)
                tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(width * 567)))
                tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])
    set_run_font(run, 12)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    footer = section.footer
    footer.paragraphs[0].text = ""
    add_page_number(footer.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    for name in ("Academic Heading 1", "Academic Heading 2", "Academic Heading 3"):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    h1 = styles["Academic Heading 1"]
    h1.base_style = normal
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    h2 = styles["Academic Heading 2"]
    h2.base_style = normal
    h2.font.name = "Times New Roman"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    h3 = styles["Academic Heading 3"]
    h3.base_style = normal
    h3.font.name = "Times New Roman"
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.first_line_indent = Cm(0)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(6)
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    if "Caption Academic" not in styles:
        styles.add_style("Caption Academic", WD_STYLE_TYPE.PARAGRAPH)
    cap = styles["Caption Academic"]
    cap.base_style = normal
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(14)
    cap.font.bold = False
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if "Appendix Code" not in styles:
        styles.add_style("Appendix Code", WD_STYLE_TYPE.PARAGRAPH)
    code = styles["Appendix Code"]
    code.base_style = normal
    code.font.name = "Courier New"
    code.font.size = Pt(8)
    code._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:cs"), "Courier New")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.left_indent = Cm(0)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)

    return doc


def add_centered(doc: Document, text: str, size=14, bold=False, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size, bold=bold)
    return p


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None):
    text = text.replace("`", "")
    if bold_prefix is not None:
        bold_prefix = bold_prefix.replace("`", "")
    p = doc.add_paragraph(style=doc.styles["Normal"])
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 14, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 14)
    else:
        run = p.add_run(text)
        set_run_font(run, 14)
    return p


def add_no_indent(doc: Document, text: str, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size, bold=bold)
    return p


def add_front_matter(doc: Document):
    add_centered(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ", 12)
    add_centered(doc, "РОССИЙСКОЙ ФЕДЕРАЦИИ", 12)
    add_centered(doc, "федеральное государственное автономное образовательное учреждение", 12)
    add_centered(doc, "высшего образования", 12)
    add_centered(doc, "«Санкт-Петербургский государственный университет", 12)
    add_centered(doc, "аэрокосмического приборостроения»", 12)
    for _ in range(4):
        doc.add_paragraph()
    add_centered(doc, "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА", 14, True)
    add_centered(doc, "ПОЯСНИТЕЛЬНАЯ ЗАПИСКА", 14, True)
    doc.add_paragraph()
    add_centered(
        doc,
        "Система с функцией обучения для распознавания жестов на основе собственного словаря",
        14,
        True,
    )
    for _ in range(5):
        doc.add_paragraph()
    add_no_indent(doc, "Студент: Давар Р. Т., гр. 4215", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_no_indent(doc, "Руководитель: к. т. н. Блюм В. С.", align=WD_ALIGN_PARAGRAPH.RIGHT)
    for _ in range(5):
        doc.add_paragraph()
    add_centered(doc, "Санкт-Петербург", 14)
    add_centered(doc, "2026", 14)
    doc.add_page_break()

    add_centered(doc, "ЗАДАНИЕ НА ВЫПУСКНУЮ КВАЛИФИКАЦИОННУЮ РАБОТУ", 14, True)
    doc.add_paragraph()
    add_body(doc, "Тема работы: система с функцией обучения для распознавания жестов на основе собственного словаря.")
    add_body(doc, "Цель работы: разработка интеллектуальной системы управления компьютером с помощью пользовательских жестов.")
    add_body(doc, "Основные задачи работы: анализ предметной области, проектирование архитектуры, разработка модуля распознавания жестов, реализация интерфейса обучения и привязки команд, тестирование разработанного приложения.")
    add_body(doc, "Примечание: данная страница является рабочей заготовкой и может быть заменена официальным заданием кафедры.")
    doc.add_page_break()

    add_centered(doc, "АННОТАЦИЯ", 14, True)
    add_body(doc, "Выпускная квалификационная работа посвящена разработке интеллектуальной системы управления компьютером с помощью жестов. Ключевой особенностью разрабатываемой системы является возможность создания пользователем собственного словаря жестов, обучения модели распознавания на пользовательских данных и привязки распознанных жестов к произвольным командам операционной системы.")
    add_body(doc, "Актуальность темы обусловлена развитием естественных интерфейсов человеко-компьютерного взаимодействия, а также потребностью в локальных, персонализируемых и доступных средствах управления компьютером без использования специализированного оборудования.")
    add_body(doc, "Целью работы является разработка desktop-приложения GestureBind, обеспечивающего запись пользовательских жестов, обучение классификатора, распознавание жестов в реальном времени, выполнение связанных команд и предоставление дружелюбного пользовательского интерфейса для настройки системы.")
    add_body(doc, "В работе рассматриваются методы компьютерного зрения, извлечения ключевых точек руки, классификации жестов на основе пользовательских данных, оптимизации гиперпараметров и проектирования интерфейсов для систем машинного обучения.")
    add_body(doc, "В результате работы планируется получить прототип приложения, пригодный для демонстрации на защите и дальнейшего развития в направлении ассистивных жестовых интерфейсов.")
    doc.add_paragraph()
    add_body(doc, "Ключевые слова: жестовое управление, компьютерное зрение, MediaPipe Hands, машинное обучение, KNN, пользовательский словарь жестов, desktop-приложение, friendly user interface.")
    doc.add_page_break()

    add_centered(doc, "СОДЕРЖАНИЕ", 14, True)
    toc_items = [
        ("ВВЕДЕНИЕ", "7"),
        ("1 Теоретические сведения и анализ предметной области", "8"),
        ("1.1 Человеко-компьютерное взаимодействие и жестовое управление", "8"),
        ("1.2 Современные подходы к распознаванию жестов", "9"),
        ("1.3 Методы детекции руки и извлечения ключевых точек", "11"),
        ("1.4 Методы классификации жестов и обучение на пользовательских данных", "14"),
        ("1.5 Обзор существующих решений и требований к дружелюбному интерфейсу", "17"),
        ("1.6 Выводы по первой главе", "21"),
        ("2 Анализ требований и проектирование системы", "23"),
        ("2.1 Постановка задачи и сценарии использования", "23"),
        ("2.2 Функциональные и нефункциональные требования", "24"),
        ("2.3 Выбор технологического стека", "26"),
        ("2.4 Общая архитектура системы GestureBind", "27"),
        ("2.5 Проектирование информационной модели и базы данных", "29"),
        ("2.6 Проектирование привязки жестов к командам", "30"),
        ("2.7 Проектирование пользовательского интерфейса и диагностики", "32"),
        ("2.8 Выводы по второй главе", "33"),
        ("3 Разработка и реализация системы GestureBind", "34"),
        ("3.1 Организация разработки и структура программных модулей", "34"),
        ("3.2 Реализация захвата видеопотока и выделения ключевых точек руки", "35"),
        ("3.3 Нормализация landmarks и формирование признакового вектора", "37"),
        ("3.4 Реализация записи пользовательских обучающих примеров", "38"),
        ("3.5 Реализация обучения классификатора жестов", "40"),
        ("3.6 Реализация онлайн-распознавания и стабилизации результата", "41"),
        ("3.7 Реализация выполнения команд по распознанным жестам", "43"),
        ("3.8 Реализация пользовательского интерфейса и прикладного контроллера", "44"),
        ("3.9 Выводы по третьей главе", "45"),
        ("4 Разработка программного приложения и пользовательского интерфейса", "47"),
        ("4.1 Назначение прикладного интерфейса", "47"),
        ("4.2 Выбор интерфейсного слоя и организация окна приложения", "48"),
        ("4.3 Прикладной контроллер и событийная модель", "49"),
        ("4.4 Главный экран распознавания", "51"),
        ("4.5 Экран обучения и экран пользовательского словаря", "52"),
        ("4.6 Экран привязки жестов к командам", "53"),
        ("4.7 Настройки, конфигурация и диагностика", "54"),
        ("4.8 Хранение состояния приложения и журналирование", "56"),
        ("4.9 Выводы по четвертой главе", "57"),
        ("5 Тестирование и оценка результатов", "58"),
        ("5.1 Методика тестирования", "58"),
        ("5.2 Автоматизированное модульное тестирование", "59"),
        ("5.3 Проверка пользовательских сценариев", "60"),
        ("5.4 Оценка качества распознавания жестов", "62"),
        ("5.5 Оценка производительности и оптимизации", "63"),
        ("5.6 Ограничения разработанного решения", "65"),
        ("5.7 Направления дальнейшего развития", "66"),
        ("5.8 Выводы по пятой главе", "67"),
        ("ЗАКЛЮЧЕНИЕ", "69"),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "72"),
    ]
    for title, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(f"{title}\t{page}")
        set_run_font(run, 13)
    doc.add_page_break()

    add_centered(doc, "УСЛОВНЫЕ ОБОЗНАЧЕНИЯ", 14, True)
    terms = [
        "CV - Computer Vision, компьютерное зрение;",
        "ML - Machine Learning, машинное обучение;",
        "HCI - Human-Computer Interaction, человеко-компьютерное взаимодействие;",
        "GUI - Graphical User Interface, графический пользовательский интерфейс;",
        "KNN - k-Nearest Neighbors, метод k-ближайших соседей;",
        "SVM - Support Vector Machine, метод опорных векторов;",
        "FPS - Frames Per Second, количество кадров в секунду;",
        "БД - база данных.",
    ]
    for idx, term in enumerate(terms, 1):
        add_no_indent(doc, f"{idx}. {term}")
    doc.add_page_break()

    add_centered(doc, "ВВЕДЕНИЕ", 14, True)
    add_body(doc, "Современные программные системы все чаще используют естественные способы взаимодействия с пользователем. Помимо клавиатуры и мыши применяются сенсорное управление, распознавание позы и жестов. Такие способы ввода позволяют сделать работу с компьютером более гибкой и доступной, особенно в сценариях презентаций, мультимедиа, hands-free управления и ассистивных интерфейсов.")
    add_body(doc, "Одним из перспективных направлений является распознавание жестов руки с использованием компьютерного зрения. При наличии обычной веб-камеры система может получать видеопоток, выделять ключевые точки руки, классифицировать жест и выполнять связанную с ним команду. При этом важной проблемой остается персонализация: заранее заданный набор жестов не всегда соответствует привычкам конкретного пользователя.")
    add_body(doc, "Целью выпускной квалификационной работы является разработка системы с функцией обучения для распознавания жестов на основе собственного словаря пользователя. Для достижения цели необходимо проанализировать предметную область, выбрать технологический стек, спроектировать архитектуру приложения, реализовать модуль распознавания жестов, разработать дружелюбный интерфейс обучения и привязки команд, а также провести тестирование разработанного решения.")
    add_body(doc, "Объектом исследования являются методы и программные средства человеко-компьютерного взаимодействия, основанные на компьютерном зрении и машинном обучении. Предметом исследования является разработка desktop-системы распознавания пользовательских жестов с механизмом обучения собственного словаря и выполнением команд компьютера.")
    doc.add_page_break()


def split_markdown_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        rows.append(parts)
    return rows


def add_markdown_table(doc: Document, table_lines: list[str]):
    rows = split_markdown_table([table_lines[0]] + table_lines[2:])
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_borders(table)
    widths_by_cols = {
        5: [3.0, 3.0, 3.8, 4.0, 3.0],
        4: [3.5, 4.5, 4.8, 4.2],
        3: [4.0, 6.0, 6.0],
    }
    widths = widths_by_cols.get(len(rows[0]), [16.5 / len(rows[0])] * len(rows[0]))
    set_table_width(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                shade_cell(cell, "EDEDED")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text.replace("`", ""))
            set_run_font(run, 10.5 if len(rows[0]) >= 4 else 11, bold=(r_idx == 0))
    doc.add_paragraph()


def add_sources(doc: Document, sources: list[str]):
    doc.add_page_break()
    add_centered(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 14, True)
    for source in sources:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(source)
        set_run_font(run, 14)


def add_conclusion_from_markdown(doc: Document, md_path: Path):
    text = md_path.read_text(encoding="utf-8")
    buffer: list[str] = []

    def flush_para():
        nonlocal buffer
        if buffer:
            paragraph = " ".join(part.strip() for part in buffer).strip()
            if paragraph:
                add_body(doc, paragraph)
            buffer = []

    doc.add_page_break()
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush_para()
            continue
        if line.startswith("# "):
            flush_para()
            add_centered(doc, line[2:].strip(), 14, True, after=12)
            continue
        buffer.append(line)
    flush_para()


def add_chapter_from_markdown(doc: Document, md_path: Path) -> list[str]:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    buffer: list[str] = []
    table_buffer: list[str] = []
    list_buffer: list[str] = []
    list_kind: str | None = None
    sources: list[str] = []
    in_sources = False
    seen_main_heading = False

    def flush_para():
        nonlocal buffer
        if buffer:
            paragraph = " ".join(part.strip() for part in buffer).strip()
            if paragraph:
                add_body(doc, paragraph)
            buffer = []

    def flush_list():
        nonlocal list_buffer, list_kind
        if not list_buffer:
            return
        text = " ".join(part.strip() for part in list_buffer).strip()
        if text:
            if list_kind == "number":
                p = doc.add_paragraph(style=doc.styles["Normal"])
                p.paragraph_format.left_indent = Cm(1.25)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text.replace("`", ""))
                set_run_font(run, 14)
            elif list_kind == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Cm(1.25)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text.replace("`", ""))
                set_run_font(run, 14)
        list_buffer = []
        list_kind = None

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            flush_list()
            flush_para()
            flush_table()
            continue
        if list_buffer and raw[:1].isspace() and not line.startswith("|"):
            list_buffer.append(line.strip())
            continue
        if list_buffer:
            flush_list()
        if line.startswith("## Источники"):
            flush_list()
            flush_para()
            flush_table()
            in_sources = True
            continue
        if in_sources:
            if re.match(r"^\d+\.\s+", line):
                sources.append(line)
            continue
        if line.startswith("|"):
            flush_list()
            flush_para()
            table_buffer.append(line)
            continue
        if line.startswith("# "):
            flush_list()
            flush_para()
            flush_table()
            if seen_main_heading:
                doc.add_page_break()
            seen_main_heading = True
            p = doc.add_paragraph(style="Academic Heading 1")
            p.paragraph_format.page_break_before = False
            run = p.add_run(line[2:].strip())
            set_run_font(run, 14, bold=True)
            continue
        if line.startswith("## "):
            flush_list()
            flush_para()
            flush_table()
            p = doc.add_paragraph(style="Academic Heading 2")
            run = p.add_run(line[3:].strip())
            set_run_font(run, 14, bold=True)
            continue
        if line.startswith("### "):
            flush_list()
            flush_para()
            flush_table()
            p = doc.add_paragraph(style="Academic Heading 3")
            run = p.add_run(line[4:].strip())
            set_run_font(run, 14, bold=True)
            continue
        if re.match(r"^\d+\.\s+", line):
            flush_para()
            flush_table()
            list_buffer = [line.strip()]
            list_kind = "number"
            continue
        if line.startswith("- "):
            flush_para()
            flush_table()
            list_buffer = [line[2:].strip()]
            list_kind = "bullet"
            continue
        if line.startswith("Таблица "):
            flush_list()
            flush_para()
            flush_table()
            p = doc.add_paragraph(style="Caption Academic")
            run = p.add_run(line)
            set_run_font(run, 14)
            continue
        if line.startswith("Рисунок "):
            flush_list()
            flush_para()
            flush_table()
            p = doc.add_paragraph(style="Caption Academic")
            run = p.add_run(line)
            set_run_font(run, 14)
            continue
        buffer.append(line)

    flush_list()
    flush_para()
    flush_table()
    return sources


def main():
    doc = setup_document()
    add_front_matter(doc)
    all_sources: list[str] = []
    for idx, chapter_md in enumerate(CHAPTER_MDS):
        if idx > 0:
            doc.add_page_break()
        all_sources.extend(add_chapter_from_markdown(doc, chapter_md))
    add_conclusion_from_markdown(doc, CONCLUSION_MD)
    add_sources(doc, all_sources)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
