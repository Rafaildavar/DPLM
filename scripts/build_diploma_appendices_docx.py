from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Приложения_к_диплому_GestureBind.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(header_cells[idx], header, bold=True)
        set_cell_shading(header_cells[idx], "D9EAF7")
        header_cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()


def add_code(doc: Document, title: str, code: str) -> None:
    p = doc.add_paragraph(title)
    p.runs[0].bold = True
    p.paragraph_format.space_after = Pt(4)
    for line in code.strip("\n").splitlines():
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.25)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(30, 30, 30)


def add_appendix_heading(doc: Document, title: str) -> None:
    if len(doc.paragraphs) > 1:
        doc.add_section(WD_SECTION_START.NEW_PAGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.49)
    p.paragraph_format.space_after = Pt(6)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.18)
    section.right_margin = Inches(0.8)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)


def main() -> None:
    doc = Document()
    configure(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ПРИЛОЖЕНИЯ К ВЫПУСКНОЙ КВАЛИФИКАЦИОННОЙ РАБОТЕ")
    run.bold = True
    run.font.size = Pt(16)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Система GestureBind для распознавания пользовательских жестов").italic = True
    add_body(
        doc,
        "Файл подготовлен как отдельный блок приложений, который можно вставить после списка использованных источников. "
        "Материалы дополняют основную часть диплома: показывают структуру проекта, ORM-модель базы данных, ключевые алгоритмы, "
        "команды запуска и проверочные метрики.",
    )

    add_appendix_heading(doc, "Приложение А. Структура программного проекта")
    add_body(
        doc,
        "В приложении приведена структура основных каталогов и модулей GestureBind. Она показывает разделение системы на слой компьютерного зрения, "
        "сервисы приложения, ORM-модель, Flet-интерфейс и автоматизированные тесты.",
    )
    add_table(
        doc,
        ["Каталог или модуль", "Назначение"],
        [
            ["app/models/database.py", "ORM-модель данных, подключение к БД, создание таблиц и сессий SQLAlchemy."],
            ["app/services/gesture_samples.py", "Синхронизация записанных sample_*.npy с таблицами gestures и gesture_samples."],
            ["app/gesture_online_infer.py", "Онлайн-инференс жестов на основе модели KNN и признаков landmarks."],
            ["app/services/gesture_command_bridge.py", "Проверка правил привязки, выбор команды для жеста и запись истории выполнения."],
            ["app/services/command_executor.py", "Исполнение действий: запуск приложений, URL, горячие клавиши, мультимедиа, скрипты."],
            ["app/flet_app/controller.py", "Прикладной контроллер GUI: камера, распознавание, обучение, команды, диагностика."],
            ["app/flet_app/views", "Экраны Flet-приложения: главная, обучение, жесты, привязки, настройки, голос."],
            ["tests/unit", "Unit-тесты сервисов, БД, привязок, команд, записи примеров, Flet-контроллера и UI-текстов."],
            ["scripts/evaluate_gesture_metrics.py", "Расчет метрик классификации на локальном датасете жестов."],
            ["docs/figures/bi", "CSV-данные для построения аналитических графиков в Tableau Public или Power BI."],
        ],
        [2.2, 4.4],
    )

    add_appendix_heading(doc, "Приложение Б. ORM-модель и структура базы данных")
    add_body(
        doc,
        "База данных хранит пользовательский словарь жестов, метаданные обучающих примеров, сведения об обученных моделях, команды, настройки и журнал распознавания. "
        "Большие массивы признаков сохраняются в файловой системе, а в БД фиксируются пути и служебные параметры.",
    )
    add_table(
        doc,
        ["Таблица", "Ключевые поля", "Назначение"],
        [
            ["users", "id, username, display_name, created_at", "Пользователи системы и персонализация словаря жестов."],
            ["gestures", "id, label, samples_path, model_class_id, accuracy, is_two_hands, is_active", "Метки жестов, состояние активности и связь с обученной моделью."],
            ["gesture_samples", "id, gesture_id, sample_index, features_path, frames, hand_count, source", "Метаданные записанных обучающих примеров gesture -> sample_*.npy."],
            ["recognition_models", "id, name, algorithm, model_path, classes_path, feature_dim, n_samples, accuracy", "Сведения о версиях обученных моделей распознавания."],
            ["commands", "id, name, platform, script_path, action_spec, gesture_id, is_active", "Команды, которые могут быть привязаны к пользовательским жестам."],
            ["settings", "key, value", "Глобальные настройки приложения в формате key-value."],
            ["gesture_history", "id, gesture_id, command_id, executed_at", "История выполненных пар gesture -> command."],
            ["recognition_logs", "id, label, confidence, model_id, gesture_id, session_id, executed", "Журнал событий распознавания и факта выполнения команды."],
            ["app_sessions", "id, user_id, app_version, platform, started_at, ended_at", "История запусков приложения и диагностика сессий."],
        ],
        [1.4, 2.5, 2.7],
    )
    add_body(
        doc,
        "Основные связи: один пользователь может иметь несколько жестов; один жест может иметь несколько обучающих примеров; одна активная команда может быть привязана к одному жесту; "
        "журнал распознавания связывает результат классификатора с моделью, жестом и сессией приложения.",
    )

    add_appendix_heading(doc, "Приложение В. Ключевые алгоритмы работы с пользовательскими жестами")
    add_body(
        doc,
        "Ниже приведены укороченные листинги, отражающие логику записи пользовательских примеров, синхронизации с базой данных и загрузки датасета для обучения. "
        "Полная версия находится в модуле app/services/gesture_samples.py.",
    )
    add_code(
        doc,
        "Листинг В.1 - Запись метаданных обучающего примера",
        """
def record_gesture_sample(session, *, label, sample_index, features_path,
                          frames=None, hand_count=None, source="camera"):
    gesture = ensure_gesture(
        session,
        label=label,
        samples_path=Path(features_path).parent,
        commit=False,
    )
    row = (
        session.query(GestureSample)
        .filter(GestureSample.gesture_id == gesture.id)
        .filter(GestureSample.sample_index == int(sample_index))
        .first()
    )
    if row is None:
        row = GestureSample(gesture_id=gesture.id, sample_index=int(sample_index))
        session.add(row)
    row.features_path = project_relative(features_path)
    row.frames = int(frames) if frames is not None else None
    row.hand_count = int(hand_count) if hand_count is not None else None
    row.source = source or "camera"
    session.commit()
    return row
""",
    )
    add_code(
        doc,
        "Листинг В.2 - Формирование признака из массива landmarks",
        """
arr = np.load(sample_path)
if arr.ndim == 3:
    t_len, points, coords = arr.shape
    arr = arr.reshape(t_len, points * coords)
feat = arr.mean(axis=0).astype(np.float32)
X.append(feat)
y.append(class_idx)
""",
    )

    add_appendix_heading(doc, "Приложение Г. Команды запуска, обучения и проверки")
    add_body(
        doc,
        "Приложение фиксирует команды, с помощью которых можно воспроизвести запуск приложения, расчет метрик и unit-тестирование. "
        "Команды выполняются из корня проекта.",
    )
    add_table(
        doc,
        ["Действие", "Команда"],
        [
            ["Запуск Flet-приложения", "python -m app.flet_app.main"],
            ["Проверка окружения", "python scripts/check_env.py"],
            ["Инициализация демонстрационных данных", "python scripts/seed_database.py"],
            ["Расчет метрик жестов", "python scripts/evaluate_gesture_metrics.py --out /tmp/dplm_metrics_current.json"],
            ["Unit-тесты стабильного набора", "pytest tests/unit --ignore=tests/unit/test_app_controller.py --no-cov"],
            ["Тесты Flet-экрана обучения", "pytest tests/unit/test_flet_training_view.py --no-cov"],
            ["Запуск GUI через shell-скрипт", "bash scripts/launch_app.sh"],
        ],
        [2.1, 4.5],
    )

    add_appendix_heading(doc, "Приложение Д. Результаты тестирования и метрики")
    add_body(
        doc,
        "В таблицах приведены данные, использованные для аналитических графиков в главе 5. Эти значения можно использовать как приложение к разделу тестирования.",
    )
    add_table(
        doc,
        ["Функциональная группа", "Пройдено", "Ошибок", "Комментарий"],
        [
            ["Конфигурация и диагностика", "10", "0", "config, pointer config, diagnostics"],
            ["Правила привязок и безопасность", "32", "0", "threshold, cooldown, dangerous actions"],
            ["Команды и системные действия", "24", "0", "command executor, DB command sync"],
            ["БД и хранение данных", "19", "0", "database models, seed, gesture samples"],
            ["CV, запись и инференс", "14", "0", "recording cycle, training, online inference"],
            ["Flet-контроллер и UI", "14", "0", "controller commands, GUI state helpers"],
            ["Управление указателем", "12", "0", "pointer movement, click, swipe"],
            ["Голосовой ассистент", "12", "0", "wake word, command processing"],
            ["Итого", "137", "0", "Все выбранные unit-тесты пройдены"],
        ],
        [2.5, 0.9, 0.8, 2.4],
    )
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Каталог датасета", "data/gestures"],
            ["Классы", "CTRLZ, New"],
            ["Количество примеров", "41"],
            ["Примеров CTRLZ", "21"],
            ["Примеров New", "20"],
            ["Размерность признаков", "42"],
            ["Форма примера", "(30, 21, 2)"],
            ["Фолдов cross-validation", "5"],
            ["Число соседей KNN", "5"],
        ],
        [2.7, 3.9],
    )
    doc.add_page_break()
    metric_caption = doc.add_paragraph("Таблица Д.3 - Сравнение вариантов обработки признаков KNN")
    metric_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metric_caption.runs[0].bold = True
    add_table(
        doc,
        ["Вариант обработки", "Accuracy", "Macro Precision", "Macro Recall", "Macro F1"],
        [
            ["Базовый pipeline", "0.9268", "0.9375", "0.9250", "0.9261"],
            ["Стандартизация", "0.9268", "0.9375", "0.9250", "0.9261"],
            ["Центрирование примера", "1.0000", "1.0000", "1.0000", "1.0000"],
            ["L2-нормализация", "0.9024", "0.9200", "0.9000", "0.9010"],
        ],
        [2.2, 1.0, 1.25, 1.15, 1.0],
    )

    add_appendix_heading(doc, "Приложение Е. Пользовательские сценарии проверки")
    add_body(
        doc,
        "Сценарии можно использовать как чек-лист ручной приемки приложения перед демонстрацией.",
    )
    add_table(
        doc,
        ["Сценарий", "Ожидаемый результат"],
        [
            ["Запуск приложения", "Открывается главное окно, отображается статус камеры, модели и БД."],
            ["Запись нового жеста", "Создаются файлы sample_*.npy, а таблицы gestures и gesture_samples обновляются."],
            ["Обучение модели", "Формируются knn.pkl, classes.json, feature_dim.txt; активные жесты доступны для привязки."],
            ["Привязка жеста к команде", "Команда сохраняется в БД, action_spec валиден, связь видна в интерфейсе."],
            ["Распознавание в реальном времени", "При выполнении жеста отображается метка, confidence и статус выполнения команды."],
            ["Защита от случайного запуска", "Команда не выполняется при низком confidence, cooldown или отключенном автоисполнении."],
            ["Удаление сэмплов", "Удаление через экран обучения удаляет метаданные и обновляет список накопленных классов."],
            ["Диагностика", "Экран настроек показывает состояние камеры, путей, модели и базы данных."],
        ],
        [2.4, 4.2],
    )

    add_appendix_heading(doc, "Приложение Ж. Перечень дополнительных иллюстраций")
    add_body(
        doc,
        "Если потребуется расширить приложения визуальными материалами, целесообразно вынести туда дополнительные скриншоты, не перегружая основную часть диплома.",
    )
    add_table(
        doc,
        ["Материал", "Источник", "Назначение"],
        [
            ["ER-диаграмма БД", "DBeaver", "Подтверждение структуры ORM и связей между таблицами."],
            ["Содержимое gestures", "DBeaver", "Демонстрация пользовательского словаря жестов."],
            ["Содержимое gesture_samples", "DBeaver", "Доказательство записи метаданных обучающих примеров."],
            ["Содержимое commands", "DBeaver", "Пример хранения пользовательских команд и action_spec."],
            ["Содержимое recognition_logs", "DBeaver", "Журнал распознавания и выполнения команд."],
            ["Папка data/gestures", "Finder/терминал", "Файловая структура sample_*.npy."],
            ["Папка models", "Finder/терминал", "Артефакты обученной модели."],
        ],
        [2.1, 1.6, 2.9],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
