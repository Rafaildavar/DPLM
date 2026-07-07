from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOC_IN = ROOT / "docs" / "Диплом все.docx"
DOC_OUT = ROOT / "docs" / "Диплом все - с графиками Tableau.docx"
FIG_DIR = ROOT / "docs" / "figures"


def copy_tableau_figures() -> dict[str, Path]:
    sources = {
        "unit_tests": Path("/Users/remi/Desktop/Снимок экрана 2026-06-06 в 18.40.31.png"),
        "dataset": Path("/Users/remi/Desktop/Снимок экрана 2026-06-06 в 18.32.40.png"),
        "metrics": FIG_DIR / "fig_5_3_tableau_metrics.png",
    }
    targets = {
        "testing_scheme": FIG_DIR / "fig_5_1_testing_scheme.png",
        "scenario_sequence": FIG_DIR / "fig_5_2_user_test_sequence.png",
        "unit_tests": FIG_DIR / "fig_5_3_tableau_unit_tests.png",
        "dataset": FIG_DIR / "fig_5_4_tableau_dataset_distribution.png",
        "metrics": FIG_DIR / "fig_5_5_tableau_metrics.png",
    }
    for key in ("unit_tests", "dataset", "metrics"):
        if not sources[key].exists():
            raise FileNotFoundError(sources[key])
        shutil.copy2(sources[key], targets[key])
    return targets


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_paragraph_startswith(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def insert_paragraph_after_table(table: Table, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_paragraph = Paragraph(new_p, table._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def add_picture_to_paragraph(paragraph: Paragraph, image_path: Path, width: float) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def insert_picture_before_caption(doc: Document, caption_text: str, image_path: Path, width: float = 6.2) -> None:
    caption = find_paragraph(doc, caption_text)
    image_paragraph = insert_paragraph_before(caption)
    add_picture_to_paragraph(image_paragraph, image_path, width)


def insert_picture_after_paragraph(
    paragraph: Paragraph,
    image_path: Path,
    caption_text: str,
    width: float = 6.2,
) -> Paragraph:
    image_paragraph = insert_paragraph_after(paragraph)
    add_picture_to_paragraph(image_paragraph, image_path, width)
    caption = insert_paragraph_after(image_paragraph, caption_text, "Caption Academic")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return caption


def insert_picture_after_table(
    table: Table,
    image_path: Path,
    caption_text: str,
    width: float = 6.2,
) -> Paragraph:
    caption = insert_paragraph_after_table(table, caption_text, "Caption Academic")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph = insert_paragraph_after_table(table)
    add_picture_to_paragraph(image_paragraph, image_path, width)
    table._tbl.addnext(image_paragraph._p)
    image_paragraph._p.addnext(caption._p)
    return caption


def replace_paragraph(doc: Document, prefix: str, text: str) -> None:
    paragraph = find_paragraph_startswith(doc, prefix)
    paragraph.text = text


def set_table_rows(table: Table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    if len(table.rows) > len(rows):
        raise ValueError(f"Unexpected row count: {len(table.rows)} != {len(rows)}")
    for row, values in zip(table.rows, rows):
        if len(row.cells) != len(values):
            raise ValueError(f"Unexpected column count: {len(row.cells)} != {len(values)}")
        for cell, value in zip(row.cells, values):
            cell.text = value


def set_cell_width(cell, width_inch: float) -> None:
    cell.width = Inches(width_inch)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_inch * 1440)))


def format_table(table: Table, widths: list[float], font_size: float = 10.0) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index == 0 or col_index > 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.bold = row_index == 0


def update_chapter_five_text(doc: Document) -> None:
    replace_paragraph(
        doc,
        "Методика тестирования включала четыре уровня.",
        "Методика тестирования включала четыре уровня, которые сведены в таблицу 5.1. Первый уровень - модульные тесты отдельных сервисов: настройки, база данных, правила привязки, выполнение команд, запись событий, подготовка данных и онлайн-инференс без руки в кадре. Второй уровень - проверка пользовательских сценариев: создание жеста, запись примеров, обучение модели, импорт жестов в БД, назначение команды и запуск распознавания. Третий уровень - оценка качества классификации на локальном наборе записанных жестов. Четвертый уровень - анализ ограничений и устойчивости решения. Общая логика проверки показана на рисунке 5.1.",
    )
    replace_paragraph(
        doc,
        "Проверка выполнялась локально на рабочем дереве проекта.",
        "Проверка выполнялась локально на рабочем дереве проекта. В качестве автоматизированного тестового фреймворка использовался pytest. Для оценки качества классификации использовался отдельный скрипт, который загружает записанные файлы sample_*.npy, формирует признаки так же, как модуль обучения, и считает accuracy, precision, recall и F1-score. Результаты скрипта выгружались в CSV и дополнительно анализировались в Tableau Public: в Tableau были получены сводные BI-графики распределения данных, тестового покрытия и метрик классификации.",
    )
    replace_paragraph(
        doc,
        "Стабильный набор тестов по ядру жестовой системы был запущен",
        "Стабильный набор тестов по ядру жестовой системы был запущен без общего coverage-gate, чтобы проверить именно функциональную корректность модулей. В этот набор вошли тесты правил привязки, БД, диагностики, сервисов команд, онлайн-инференса, сэмплов жестов, управления указателем, событий распознавания, записи примеров, обучения классификатора, Flet-контроллера, UI-текстов и голосового ассистента. Результат по функциональным группам приведен в таблице 5.2, а параметры прогона - в таблице 5.3.",
    )
    replace_paragraph(
        doc,
        "Дополнительно был выполнен более широкий прогон unit-тестов",
        "Для наглядного анализа данные таблиц 5.2 и 5.3 были выгружены в Tableau Public и представлены в виде горизонтальной диаграммы. На рисунке 5.3 показано распределение 137 успешно пройденных unit-тестов по функциональным группам. Наибольшее число проверок относится к правилам привязок и безопасности, командам и системным действиям, а также хранению данных.",
    )
    replace_paragraph(
        doc,
        "Полученный результат показывает, что ядро жестовой системы покрыто",
        "Полученный результат показывает, что ядро жестовой системы покрыто автоматизированными проверками по основным рисковым зонам: правила привязок, БД, команды, запись обучающих примеров и поведение Flet-контроллера. Поэтому дальнейшее расширение тестов целесообразно направить прежде всего на UI-сценарии с камерой и интеграционные проверки полного пользовательского цикла.",
    )
    replace_paragraph(
        doc,
        "Помимо unit-тестов была сформирована таблица сценариев",
        "Помимо unit-тестов была сформирована таблица 5.4 со сценариями, которые должен выполнить пользователь в рабочем приложении. Эти сценарии соответствуют функциональным требованиям, сформулированным во второй главе: запись жестов, обучение модели, распознавание в реальном времени, привязка команд и диагностика. Последовательность пользовательской проверки представлена на рисунке 5.2.",
    )
    replace_paragraph(
        doc,
        "Второй сценарий - создание пользовательского жеста.",
        "Второй сценарий - создание пользовательского жеста. Пользователь задает метку и число примеров. После записи в файловой структуре должны появиться файлы sample_*.npy, а экран обучения должен показать сообщения процесса.",
    )
    replace_paragraph(
        doc,
        "Качество распознавания оценивалось на локальном наборе записанных примеров.",
        "Качество распознавания оценивалось на локальном наборе записанных примеров. Оценочный скрипт использует тот же базовый pipeline признаков, что и модуль обучения: загрузка файлов sample_*.npy, преобразование последовательности landmarks в двумерный массив, усреднение по времени, приведение к float32 и выравнивание размерности признаков. Параметры оценочного набора приведены в таблице 5.5, а распределение обучающих примеров по классам было получено в Tableau Public и показано на рисунке 5.4.",
    )
    replace_paragraph(
        doc,
        "В последнем локальном замере было прочитано 41 обучающий пример",
        "В актуальном локальном замере было прочитано 41 обучающий пример из двух классов. Класс CTRLZ содержал 21 пример, класс New - 20 примеров. Размерность признакового вектора после выравнивания составила 42, то есть использовался одноручный формат признаков: 21 ключевая точка руки и две координаты для каждой точки. Пропущенных файлов не было, что подтверждает корректность записи и чтения текущего набора sample_*.npy.",
    )
    replace_paragraph(
        doc,
        "Для оценки использовалась 5-fold cross-validation.",
        "Для оценки использовалась 5-fold cross-validation. Были рассчитаны accuracy, macro precision, macro recall и macro F1. Также сравнивались несколько вариантов обработки признаков: базовый pipeline, стандартизация на обучающей выборке, центрирование относительно среднего значения примера и L2-нормализация примера. Численные значения приведены в таблице 5.6, а тепловая карта метрик, полученная в Tableau Public, представлена на рисунке 5.5.",
    )
    replace_paragraph(
        doc,
        "Полученные значения показывают, что на доступной части текущего локального датасета два класса разделяются без ошибок.",
        "Полученные значения показывают, что базовый pipeline и стандартизация дают accuracy 0.9268 и macro F1 0.9261, L2-нормализация снижает accuracy до 0.9024, а центрирование примера дает 1.0000 по всем основным метрикам на текущем малом наборе данных. Это подтверждает важность предварительной обработки landmarks. Однако результат нельзя трактовать как окончательное доказательство устойчивости в произвольных условиях: набор данных мал, содержит только два класса и был записан в ограниченных условиях. При расширении пользовательского словаря, изменении освещения, расстояния до камеры или появлении похожих жестов метрики могут снизиться.",
    )
    replace_paragraph(
        doc,
        "В рамках отдельного эксперимента по оптимизации гиперпараметров рассматривался",
        "В текущей проверке сравнение выполнялось не через полный автоматический подбор гиперпараметров, а через несколько воспроизводимых вариантов предварительной обработки признаков для KNN-классификатора. Такой подход позволяет показать практический эффект подготовки landmarks и при этом не заявлять результаты, которые не подтверждены текущим запуском.",
    )
    replace_paragraph(
        doc,
        "В экспериментальном отчете по оптимизации приведены следующие результаты",
        "Ранее подготовленная сводка оптимизации была проверена отдельно и не включена в итоговую оценку как доказательный результат, поскольку в текущем рабочем окружении она не подтверждается воспроизводимым запуском: скрипт оптимизации требует пакет Optuna, отсутствующий в окружении, а отдельный замер задержки инференса в проекте не представлен. Поэтому в итоговой версии раздела использованы только метрики, полученные воспроизводимым скриптом evaluate_gesture_metrics.py на текущем локальном датасете. Проверенная сводка приведена в таблице 5.7.",
    )
    replace_paragraph(
        doc,
        "В другом эксперименте с текущим набором данных лучшая конфигурация дала значение",
        "На текущем наборе данных базовый pipeline KNN при k = 5 дал accuracy 0.9268 и macro F1 0.9261. Лучший проверенный вариант предварительной обработки - центрирование признаков каждого примера - дал значение 1.0000 по accuracy и macro F1. Такой результат ожидаем для малого набора хорошо разделимых классов, но при расширении датасета параметры и способ обработки необходимо проверять повторно.",
    )
    replace_paragraph(
        doc,
        "Практический вывод состоит в том, что механизм оптимизации целесообразно использовать",
        "Практический вывод состоит в том, что подбор параметров и предварительной обработки целесообразно рассматривать как дополнительный этап после накопления достаточного числа примеров. Для маленького датасета результат может быть слишком оптимистичным, поэтому в работе не заявляется неподтвержденное сокращение времени инференса. Основной подтвержденный эффект состоит в том, что вариант обработки признаков влияет на качество KNN и должен подбираться на валидации.",
    )


def update_chapter_five_tables(doc: Document) -> None:
    find_paragraph(doc, "Таблица 5.3 - Результат расширенного предварительного прогона").text = (
        "Таблица 5.3 - Параметры стабильного прогона unit-тестов"
    )
    set_table_rows(
        doc.tables[26],
        [
            ["Группа тестов", "Количество", "Результат"],
            ["Конфигурация и диагностика", "10", "Пройдено"],
            ["Правила привязок и безопасность", "32", "Пройдено"],
            ["Команды и системные действия", "24", "Пройдено"],
            ["БД и хранение данных", "19", "Пройдено"],
            ["CV, запись и инференс", "14", "Пройдено"],
            ["Flet-контроллер и UI", "14", "Пройдено"],
            ["Управление указателем", "12", "Пройдено"],
            ["Голосовой ассистент", "12", "Пройдено"],
            ["Итого", "137", "137 пройдено"],
        ],
    )
    set_table_rows(
        doc.tables[27],
        [
            ["Показатель", "Значение", "Интерпретация"],
            ["Команда запуска", "pytest tests/unit --ignore=tests/unit/test_app_controller.py --no-cov", "Проверка стабильного набора unit-тестов"],
            ["Успешно завершено", "137", "Все выбранные проверки пройдены"],
            ["Завершилось ошибкой", "0", "Блокирующих ошибок не выявлено"],
            ["Исключено", "test_app_controller.py", "Старый PySide/QML-контроллер не относится к текущей Flet-версии"],
            ["Coverage", "не считался", "Цель прогона - функциональная проверка, а не покрытие строк"],
        ],
    )
    set_table_rows(
        doc.tables[29],
        [
            ["Параметр", "Значение"],
            ["Каталог данных", "data/gestures"],
            ["Число доступных примеров", "41"],
            ["Количество классов", "2"],
            ["Классы", "CTRLZ, New"],
            ["Примеров класса CTRLZ", "21"],
            ["Примеров класса New", "20"],
            ["Пропущено файлов", "0"],
            ["Размерность признаков", "42"],
            ["Число фолдов", "5"],
            ["Число соседей KNN", "5"],
        ],
    )
    set_table_rows(
        doc.tables[30],
        [
            ["Вариант обработки", "Accuracy", "Macro precision", "Macro recall", "Macro F1"],
            ["Базовый pipeline", "0.9268", "0.9375", "0.9250", "0.9261"],
            ["Стандартизация", "0.9268", "0.9375", "0.9250", "0.9261"],
            ["Центрирование примера", "1.0000", "1.0000", "1.0000", "1.0000"],
            ["L2-нормализация", "0.9024", "0.9200", "0.9000", "0.9010"],
        ],
    )
    find_paragraph(doc, "Таблица 5.7 - Результаты эксперимента по оптимизации KNN").text = (
        "Таблица 5.7 - Проверенная сводка сравнения вариантов KNN на текущем датасете"
    )
    set_table_rows(
        doc.tables[31],
        [
            ["Показатель", "Базовый pipeline KNN", "Лучший проверенный вариант", "Изменение / вывод"],
            ["Accuracy", "0.9268", "1.0000 (центрирование примера)", "+0.0732"],
            ["Macro F1", "0.9261", "1.0000 (центрирование примера)", "+0.0739"],
            ["Среднее время инференса", "не измерялось", "не измерялось", "не заявляется в итоговом выводе"],
        ],
    )


def update_figures(doc: Document, figures: dict[str, Path]) -> None:
    insert_picture_before_caption(
        doc,
        "Рисунок 5.1 - Общая схема тестирования GestureBind",
        figures["testing_scheme"],
        5.8,
    )
    insert_picture_before_caption(
        doc,
        "Рисунок 5.2 - Последовательность пользовательского тестового сценария",
        figures["scenario_sequence"],
        5.8,
    )

    unit_anchor = find_paragraph_startswith(doc, "Полученный результат показывает, что ядро жестовой системы покрыто")
    insert_picture_after_paragraph(
        unit_anchor,
        figures["unit_tests"],
        "Рисунок 5.3 - Распределение пройденных unit-тестов по функциональным группам, полученное в Tableau Public",
        6.2,
    )

    dataset_caption_anchor = doc.tables[29]
    insert_picture_after_table(
        dataset_caption_anchor,
        figures["dataset"],
        "Рисунок 5.4 - Распределение обучающих примеров по классам локального датасета, полученное в Tableau Public",
        6.2,
    )

    old_metrics_caption = find_paragraph(doc, "Рисунок 5.3 - Сравнение вариантов обработки признаков")
    old_metrics_caption.text = "Рисунок 5.5 - Сравнение вариантов обработки признаков по метрикам классификации в Tableau Public"
    insert_picture_before_caption(
        doc,
        "Рисунок 5.5 - Сравнение вариантов обработки признаков по метрикам классификации в Tableau Public",
        figures["metrics"],
        6.2,
    )


def main() -> None:
    figures = copy_tableau_figures()
    doc = Document(DOC_IN)
    update_chapter_five_text(doc)
    update_chapter_five_tables(doc)
    update_figures(doc, figures)
    doc.save(DOC_OUT)
    print(DOC_OUT)


if __name__ == "__main__":
    main()
